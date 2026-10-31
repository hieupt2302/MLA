"""Build the Week 11 DOCX and HTML from Markdown; never rewrite the source."""
import html
import re
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "RESEARCH_REPORT_WEEK11_DISCOURSE_EFFECTIVENESS_CLC02_GROUP06.md"
OUTPUT = ROOT / "Week11_Group06_CLC02_Report"
HEADING = re.compile(r"^(#{1,3})\s+(.+)$")
LIST = re.compile(r"^\s*(?:([-*])\s+|(\d+)\.\s+)(.*)$")
INLINE = re.compile(r"`([^`]+)`|\[([^\]]+)\]\(([^\s)]+)\)|"
                    r"(https?://[^\s<>`]+)|\*\*(.+?)\*\*|\*([^*]+)\*")


def cells(line):
    return [s.strip() for s in line.strip().strip("|").split("|")]


def separator(line):
    return "|" in line and all(re.fullmatch(r":?-{3,}:?", s) for s in cells(line))


def parse_markdown(text):
    """Yield (kind, data) blocks for the report's small Markdown subset."""
    lines = text.expandtabs(4).splitlines()
    i = 0

    def starts_block(pos):
        s = lines[pos].strip()
        return (not s or HEADING.match(s) or LIST.match(s) or s == "---"
                or (pos + 1 < len(lines) and "|" in s and separator(lines[pos + 1])))

    def joined(parts):
        return "".join(s.strip() + ("\n" if s.endswith("  ") else " ")
                       for s in parts).rstrip()

    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
        elif i + 1 < len(lines) and "|" in line and separator(lines[i + 1]):
            rows = [cells(line)]
            width = len(rows[0])
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                row = cells(lines[i])
                if len(row) != width:
                    print(f"WARNING: line {i + 1}: table row has {len(row)} cells; "
                          f"expected {width}; padding/truncating.")
                rows.append((row + [""] * width)[:width])
                i += 1
            yield "table", rows
        elif HEADING.match(line):
            match = HEADING.match(line)
            yield "heading", (len(match[1]), match[2])
            i += 1
        elif line == "---":
            yield "rule", None
            i += 1
        elif LIST.match(line):
            match = LIST.match(line)
            parts = [match[3] + ("  " if lines[i].endswith("  ") else "")]
            number = match[2]
            i += 1
            while i < len(lines) and not starts_block(i) and lines[i].startswith("  "):
                parts.append(lines[i])
                i += 1
            yield "item", (number, joined(parts))
        else:
            parts = [lines[i]]
            i += 1
            while i < len(lines) and not starts_block(i):
                parts.append(lines[i])
                i += 1
            yield "paragraph", joined(parts)


def inline(text, paragraph, bold=False, italic=False):
    """Emit Word runs and return escaped HTML from the same inline tokens."""
    result = []

    def plain(value, code=False):
        run = paragraph.add_run(value)
        run.bold, run.italic = bold, italic
        if code:
            run.font.name = "Consolas"
        escaped = html.escape(value).replace("\n", "<br>\n")
        return f"<code>{escaped}</code>" if code else escaped

    end = 0
    for match in INLINE.finditer(text):
        result.append(plain(text[end:match.start()]))
        code, label, url, bare, strong, emphasis = match.groups()
        if code is not None:
            result.append(plain(code, code=True))
        elif url or bare:
            suffix = ""
            if bare:
                url = bare.rstrip(".,;:!?)")
                suffix = bare[len(url):]
                label = url
            # Only make safe web/mail links clickable in HTML.
            safe = url.lower().startswith(("http://", "https://", "mailto:"))
            body = plain(label)
            if label != url:
                plain(f" ({url})")
            result.append(f'<a href="{html.escape(url, quote=True)}">{body}</a>'
                          if safe else body + " (" + html.escape(url) + ")")
            result.append(plain(suffix))
        else:
            tag = "strong" if strong is not None else "em"
            body = inline(strong if strong is not None else emphasis, paragraph,
                          bold or strong is not None, italic or emphasis is not None)
            result.append(f"<{tag}>{body}</{tag}>")
        end = match.end()
    result.append(plain(text[end:]))
    return "".join(result)


CSS = """
body { max-width:1100px; margin:36px auto; padding:0 24px;
       font:16px/1.55 Calibri,Arial,sans-serif; color:#202a35; }
h1,h2,h3 { line-height:1.25; color:#183d58; margin-top:1.4em; }
h1 { font-size:1.85em; } h2 { font-size:1.4em; } h3 { font-size:1.15em; }
.table-scroll { overflow-x:auto; margin:18px 0; }
table { border-collapse:collapse; width:100%; font-size:14px; }
th,td { border:1px solid #aab6c2; padding:7px 9px; vertical-align:top; }
th { background:#e8eef4; text-align:left; }
.wide table { min-width:1000px; }
code { font-family:Consolas,monospace; overflow-wrap:anywhere; }
a { color:#145c9e; overflow-wrap:anywhere; } hr { border:0; border-top:1px solid #aab6c2; }
"""


def build(blocks):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name in ["Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet"]:
        doc.styles[name].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for level, size in [(1, 18), (2, 14), (3, 12)]:
        doc.styles[f"Heading {level}"].font.size = Pt(size)
    body = []
    list_tag = None
    for kind, data in blocks:
        next_tag = ("ol" if data[0] else "ul") if kind == "item" else None
        if list_tag != next_tag:
            if list_tag:
                body.append(f"</{list_tag}>")
            if next_tag:
                start = f' start="{int(data[0])}"' if next_tag == "ol" else ""
                body.append(f"<{next_tag}{start}>")
            list_tag = next_tag
        if kind == "table":
            width = len(data[0])
            table = doc.add_table(rows=0, cols=width)
            table.style = "Table Grid"
            table.autofit = True
            wide = " wide" if width >= 8 else ""
            body.append(f'<div class="table-scroll{wide}"><table>')
            for index, values in enumerate(data):
                row = table.add_row()
                if index == 0:
                    repeat = OxmlElement("w:tblHeader")
                    row._tr.get_or_add_trPr().append(repeat)
                body.append("<thead><tr>" if index == 0 else "<tr>")
                for cell, value in zip(row.cells, values):
                    p = cell.paragraphs[0]
                    rendered = inline(value, p, bold=index == 0)
                    for run in p.runs:
                        run.font.size = Pt(7.5 if width >= 8 else 8.5)
                    p.paragraph_format.space_after = Pt(2)
                    if index == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "E8EEF4")
                        cell._tc.get_or_add_tcPr().append(shade)
                    tag = "th" if index == 0 else "td"
                    body.append(f"<{tag}>{rendered}</{tag}>")
                body.append("</tr></thead><tbody>" if index == 0 else "</tr>")
            body.append("</tbody></table></div>")
        elif kind == "rule":
            p = doc.add_paragraph()
            borders = OxmlElement("w:pBdr")
            border = OxmlElement("w:bottom")
            for key, value in [("val", "single"), ("sz", "4"), ("color", "AAB6C2")]:
                border.set(qn("w:" + key), value)
            borders.append(border)
            p._p.get_or_add_pPr().append(borders)
            body.append("<hr>")
        else:
            if kind == "heading":
                level, text = data
                p, tag = doc.add_heading(level=level), f"h{level}"
            elif kind == "item":
                number, text = data
                p = doc.add_paragraph(style="Normal" if number else "List Bullet")
                if number:
                    p.paragraph_format.left_indent = Cm(0.63)
                    p.paragraph_format.first_line_indent = Cm(-0.63)
                    p.add_run(number + ". ")
                tag = "li"
            else:
                text, p, tag = data, doc.add_paragraph(), "p"
            value = inline(text, p)
            number_attr = f' value="{int(data[0])}"' if kind == "item" and data[0] else ""
            body.append(f"<{tag}{number_attr}>{value}</{tag}>")
    if list_tag:
        body.append(f"</{list_tag}>")
    page = ('<!doctype html>\n<html lang="en"><head><meta charset="utf-8">'
            '<meta name="viewport" content="width=device-width, initial-scale=1">'
            f'<title>{html.escape(OUTPUT.stem)}</title><style>{CSS}</style></head><body>\n'
            + "\n".join(body) + "\n</body></html>\n")
    return doc, page


def verify(blocks, path, page):
    doc = Document(path)
    tables = [data for kind, data in blocks if kind == "table"]
    assert len(doc.tables) == len(tables), "DOCX table count mismatch"
    for table, rows in zip(doc.tables, tables):
        assert len(table.rows) == len(rows), "DOCX row count mismatch"
        assert all(len(row.cells) == len(rows[0]) for row in table.rows), "Ragged DOCX table"
    assert "\t" not in "".join(doc.element.itertext()), "TAB in DOCX text"
    assert not doc.element.xpath(".//w:tab"), "TAB element in DOCX"
    assert "\t" not in page, "TAB in HTML"

    class TableCounter(HTMLParser):
        count = 0

        def handle_starttag(self, tag, attrs):
            if tag == "table":
                self.count += 1

    counter = TableCounter()
    counter.feed(page)
    assert counter.count == len(tables), "HTML table count mismatch"
    print(f"Verified: Markdown/DOCX/HTML tables = {len(tables)}/{len(doc.tables)}/{counter.count}; "
          "all table rows rectangular; no TABs in DOCX or HTML.")


def main():
    blocks = list(parse_markdown(SOURCE.read_text(encoding="utf-8-sig")))
    doc, page = build(blocks)
    doc.save(OUTPUT.with_suffix(".docx"))
    OUTPUT.with_suffix(".html").write_text(page, encoding="utf-8")
    verify(blocks, OUTPUT.with_suffix(".docx"), OUTPUT.with_suffix(".html").read_text(encoding="utf-8"))
    print(f"Wrote {OUTPUT.name}.docx and {OUTPUT.name}.html")


if __name__ == "__main__":
    main()
